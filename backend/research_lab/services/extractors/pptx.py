from __future__ import annotations

from typing import BinaryIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from research_lab.services.extractors.base import ParsedPart, clean_text


def _shape_texts(shape) -> list[str]:
    texts: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            texts.extend(_shape_texts(child))
    elif getattr(shape, 'has_table', False) and shape.has_table:
        for row in shape.table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            texts.append(' | '.join(c for c in cells if c))
    elif getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
        text = clean_text(shape.text_frame.text)
        if text:
            texts.append(text)
    return texts


def extract_pptx(fh: BinaryIO) -> list[ParsedPart]:
    presentation = Presentation(fh)
    parts: list[ParsedPart] = []
    for i, slide in enumerate(presentation.slides):
        title = ''
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = clean_text(slide.shapes.title.text_frame.text)
        blocks: list[str] = []
        for shape in slide.shapes:
            blocks.extend(_shape_texts(shape))
        metadata: dict = {}
        if slide.has_notes_slide:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                metadata['notes'] = notes
                blocks.append(f'[Speaker notes] {notes}')
        parts.append(ParsedPart(
            part_type='SLIDE', index=i, title=title[:500], text='\n'.join(blocks),
            slide_number=i + 1, metadata=metadata,
        ))
    return parts
