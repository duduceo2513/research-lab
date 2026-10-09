"""Deterministic, code-only rendering of a lecture's structured content. No LLM call lives here —
the Lecture Builder agent writes the content once; everything below is pure formatting, so DOCX
(and later PPTX) is an export layer on top of ``content_json``, not a second generation pass."""
from __future__ import annotations

import io


def render_plain_text(*, title: str, learning_objectives: list, content: dict) -> str:
    """Markdown-flavoured plain text — what the "Lecture" tab shows, and what the Copy/Export
    Markdown/Export TXT actions hand the user, unchanged."""
    lines = [f'# {title}' if title else '# Untitled lecture', '']

    if learning_objectives:
        lines.append('## Learning objectives')
        lines += [f'- {objective}' for objective in learning_objectives]
        lines.append('')

    for section in content.get('sections') or []:
        lines.append(f"## {section.get('title', '')}")
        if section.get('content'):
            lines.append(section['content'])
            lines.append('')
        if section.get('key_points'):
            lines.append('**Key points:**')
            lines += [f'- {point}' for point in section['key_points']]
            lines.append('')
        if section.get('examples'):
            lines.append('**Examples:**')
            lines += [f'- {example}' for example in section['examples']]
            lines.append('')
        if section.get('source_ids'):
            lines.append('_Sources: ' + ', '.join(f'[{sid}]' for sid in section['source_ids']) + '_')
            lines.append('')

    if content.get('summary'):
        lines.append('## Summary')
        lines.append(content['summary'])
        lines.append('')

    if content.get('self_check_questions'):
        lines.append('## Self-check questions')
        lines += [f'{i}. {q}' for i, q in enumerate(content['self_check_questions'], start=1)]
        lines.append('')

    if content.get('references'):
        lines.append('## References')
        for ref in content['references']:
            citation = ref.get('citation', str(ref)) if isinstance(ref, dict) else str(ref)
            lines.append(f'- {citation}')
        lines.append('')

    if content.get('missing_topics'):
        lines.append('## Topics needing more research')
        lines += [f'- {topic}' for topic in content['missing_topics']]
        lines.append('')

    return '\n'.join(lines).strip()


def render_docx(*, title: str, learning_objectives: list, content: dict) -> bytes:
    """Builds a real .docx from the same structured content ``render_plain_text`` uses — the
    export layer the spec calls for, reusing ``python-docx`` (already a dependency for reading
    downloaded .docx files)."""
    from docx import Document

    document = Document()
    document.add_heading(title or 'Untitled lecture', level=0)

    if learning_objectives:
        document.add_heading('Learning objectives', level=1)
        for objective in learning_objectives:
            document.add_paragraph(str(objective), style='List Bullet')

    for section in content.get('sections') or []:
        document.add_heading(section.get('title', ''), level=1)
        if section.get('content'):
            document.add_paragraph(section['content'])
        if section.get('key_points'):
            document.add_paragraph().add_run('Key points:').bold = True
            for point in section['key_points']:
                document.add_paragraph(str(point), style='List Bullet')
        if section.get('examples'):
            document.add_paragraph().add_run('Examples:').bold = True
            for example in section['examples']:
                document.add_paragraph(str(example), style='List Bullet')
        if section.get('source_ids'):
            note = 'Sources: ' + ', '.join(f'[{sid}]' for sid in section['source_ids'])
            document.add_paragraph().add_run(note).italic = True

    if content.get('summary'):
        document.add_heading('Summary', level=1)
        document.add_paragraph(content['summary'])

    if content.get('self_check_questions'):
        document.add_heading('Self-check questions', level=1)
        for question in content['self_check_questions']:
            document.add_paragraph(str(question), style='List Number')

    if content.get('references'):
        document.add_heading('References', level=1)
        for ref in content['references']:
            citation = ref.get('citation', str(ref)) if isinstance(ref, dict) else str(ref)
            document.add_paragraph(citation, style='List Bullet')

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
